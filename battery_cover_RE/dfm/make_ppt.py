"""12-slide DFM deck (python-pptx) from the run_dfm.py summary."""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

DARK = RGBColor(0x1F, 0x2A, 0x36)
STEEL = RGBColor(0x4C, 0x78, 0xA8)
AMBER = RGBColor(0xF2, 0x8E, 0x2B)
RED = RGBColor(0xD6, 0x27, 0x28)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
GREY = RGBColor(0x5B, 0x67, 0x70)
LIGHT = RGBColor(0xF3, 0xF5, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEAD = "Cambria"
BODY = "Calibri"
W, H = 13.333, 7.5
FOOT = "ANRMTPT0003F Battery cover  ·  ABS  ·  DFM review  ·  Rev 01"


# ============================================================ primitives
def text(slide, x, y, w, h, runs, size=14, color=DARK, font=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False):
    """runs: str | list of paragraphs; a paragraph is str or list of (text, {overrides})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for seg, over in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = seg
            f = r.font
            f.name = over.get("font", font)
            f.size = Pt(over.get("size", size))
            f.bold = over.get("bold", bold)
            f.italic = over.get("italic", italic)
            f.color.rgb = over.get("color", color)
    return tb


def bullets(slide, x, y, w, h, items, size=13, color=DARK, gap=5):
    """items: str or (bold_lead, rest)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        lead, rest = (it, "") if isinstance(it, str) else it
        r0 = p.add_run()
        r0.text = "▸  "
        r0.font.size = Pt(size)
        r0.font.color.rgb = AMBER
        r0.font.name = BODY
        r1 = p.add_run()
        r1.text = lead
        r1.font.size = Pt(size)
        r1.font.bold = bool(rest)
        r1.font.color.rgb = color
        r1.font.name = BODY
        if rest:
            r2 = p.add_run()
            r2.text = " " + rest
            r2.font.size = Pt(size)
            r2.font.color.rgb = color
            r2.font.name = BODY
    return tb


def card(slide, x, y, w, h, fill=LIGHT):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.06
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def image(slide, path, x, y, w, h):
    """Fit an image inside the box, centred, keeping aspect."""
    with Image.open(path) as im:
        iw, ih = im.size
    s = min(w / iw, h / ih)
    dw, dh = iw * s, ih * s
    return slide.shapes.add_picture(str(path), Inches(x + (w - dw) / 2), Inches(y + (h - dh) / 2), Inches(dw),
                                    Inches(dh))


def table(slide, x, y, w, rows, col_w=None, size=11, row_h=0.3, head_fill=DARK, zebra=True, bold_first_col=False,
          colors=None):
    n_r, n_c = len(rows), len(rows[0])
    shp = slide.shapes.add_table(n_r, n_c, Inches(x), Inches(y), Inches(w), Inches(row_h * n_r))
    tbl = shp.table
    col_w = col_w or [w / n_c] * n_c
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(val)
            r.font.name = BODY
            r.font.size = Pt(size)
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = head_fill
                r.font.bold = True
                r.font.color.rgb = WHITE
            else:
                cell.fill.fore_color.rgb = LIGHT if (zebra and i % 2 == 0) else WHITE
                r.font.color.rgb = DARK
                r.font.bold = bold_first_col and j == 0
                if colors and (i, j) in colors:
                    r.font.color.rgb = colors[(i, j)]
                    r.font.bold = True
    return tbl


def stat(slide, x, y, w, value, label, color=AMBER, vsize=30):
    text(slide, x, y, w, 0.6, value, size=vsize, color=color, font=HEAD, bold=True)
    text(slide, x, y + 0.62, w, 0.5, label, size=11, color=GREY)


def frame(prs, n, title, key, tag):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    text(s, 0.5, 0.32, 10.6, 0.62, title, size=28, font=HEAD, bold=True, color=DARK)
    text(s, 0.5, 0.95, 11.2, 0.45, key, size=14, color=GREY, italic=True)
    t = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(W - 2.3), Inches(0.42), Inches(1.8), Inches(0.42))
    t.adjustments[0] = 0.5
    t.fill.solid()
    t.fill.fore_color.rgb = AMBER
    t.line.fill.background()
    tf = t.text_frame
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = f"{n:02d}  {tag}"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.name = BODY
    r.font.color.rgb = WHITE
    text(s, 0.5, 7.05, 9, 0.3, FOOT, size=9, color=GREY)
    text(s, W - 1.5, 7.05, 1.0, 0.3, f"{n} / 12", size=9, color=GREY, align=PP_ALIGN.RIGHT)
    return s


def style_chart(ch, size=10):
    ch.font.size = Pt(size)
    ch.font.name = BODY
    ch.font.color.rgb = DARK


# ============================================================ deck
def build(S, repo: Path, out: Path):
    repo = Path(repo)
    F = {k: repo / v for k, v in S["figures"].items()}
    inp, dfm = S["input"], S["dfm"]
    mat = S["material"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)

    # ---------------------------------------------------------------- 1 info
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = DARK
    text(s, 0.6, 0.45, 12.2, 0.8, "Battery Cover  ANRMTPT0003F", size=36, font=HEAD, bold=True, color=WHITE)
    text(s, 0.6, 1.22, 12.2, 0.5, "DFM review for injection moulding  ·  ABS", size=20, color=AMBER)
    text(s, 0.6, 1.74, 12.2, 0.4, "TE-12XCME remote control  ·  assembly ANRMTAS0001F (item 6)  ·  PhotonX R&D",
         size=13, color=RGBColor(0xC9, 0xD1, 0xD9))
    card(s, 0.6, 2.35, 2.35, 4.4, fill=WHITE)
    card(s, 3.1, 2.35, 2.35, 4.4, fill=WHITE)
    image(s, F["r_input_outer"], 0.65, 2.4, 2.25, 4.3)
    image(s, F["r_input_inner"], 3.15, 2.4, 2.25, 4.3)
    text(s, 0.6, 6.8, 4.85, 0.3, "As received: outer (cosmetic) face  |  inner face", size=10,
         color=RGBColor(0xC9, 0xD1, 0xD9), align=PP_ALIGN.CENTER)
    sz = inp["size_xyz"]
    stats = [(f"{sz[2]:g} × {sz[0]:g} × {sz[1]:g}", "envelope L × W × H [mm]"),
             (f"{inp['volume_mm3'] / 1000:.2f} cm³", f"volume  ·  {inp['mass_g']:.2f} g in ABS"),
             (f"{inp['reuse_checks_dfm']['projected_area_Y_cm2']:.1f} cm²", "projected area (pull ±Y)"),
             ("2.0 mm", "nominal wall")]
    for i, (v, lab) in enumerate(stats):
        stat(s, 5.85 + (i % 2) * 3.55, 2.35 + (i // 2) * 1.3, 3.4, v, lab, vsize=24)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xC9, 0xD1, 0xD9)
    for shp in list(s.shapes)[-8:]:
        for p in shp.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size and r.font.size <= Pt(12):
                    r.font.color.rgb = RGBColor(0xC9, 0xD1, 0xD9)
    if "assembly" in F:
        card(s, 5.85, 5.0, 3.7, 1.95, fill=WHITE)
        image(s, F["assembly"], 5.9, 5.05, 3.6, 1.85)
    text(s, 9.75, 5.0, 3.1, 1.95, [
        [("Found in the as-received model", {"bold": True, "color": WHITE, "size": 13})],
        [(f"1,372 mm² of walls at 0° draft", {"color": AMBER})],
        [(f"{inp['undercut_area_mm2']:.0f} mm² undercut (hook gap)", {"color": AMBER})],
        [("sharp internal corners (R0)", {"color": AMBER})],
        [("Draft + radii fixed in part_DFM.step; the hook gap is moulded by a lifter (the cover stays closed)",
          {"color": WHITE, "bold": True})],
    ], size=12, color=WHITE)
    s.notes_slide.notes_text_frame.text = (
        "Input: input/part.step = design-intent battery cover rebuilt from the FreeCAD mesh (repo battery_cover_RE). "
        "Part identity from assembly drawing ANRMTAS0001F sheet 2 BOM item 6 (71.5 x 31.0 x 9.0 envelope). "
        "Material ABS general purpose, typical datasheet values used for all calculations; confirm with the chosen "
        "grade. All geometry numbers are measured from the STEP files by battery_cover_RE/dfm/run_dfm.py.")

    # ---------------------------------------------------------------- 2 PL / PS
    s = frame(prs, 2, "Parting line & parting surface",
              "Flat PL on the outer-face edge; one 7 mm step around the snap latch. One lifter per cavity for the hook.",
              "PL / PS")
    image(s, F["parting"], 0.4, 1.5, 8.3, 3.05)
    image(s, F["ps_section"], 0.4, 4.6, 8.3, 2.35)
    card(s, 8.95, 1.55, 3.9, 5.35)
    bullets(s, 9.1, 1.7, 3.65, 5.1, [
        ("Pull ±Y", "(plate normal). Part stays on the core (B) and is ejected from B."),
        ("A / cavity:", "cosmetic outer face, bumps, thumb recess, latch catch faces (blue)."),
        ("B / core:", "ribs, rails, lip, hook, latch inner face (orange)."),
        ("Main PL:", "outer-face perimeter at Y 20 (red). The witness line sits on the hidden edge, not on the cosmetic face."),
        ("Stepped PS:", "drops 7 mm to Y 13 around the snap arm. Tool shut-off faces ≥ 5°."),
        ("Hook gap (purple):", "moulded by a 10° angled lifter on the B side. No hole in the cover."),
        (f"{S['parting']['pl_edges']} PL edges", "found automatically at the A/B face boundary."),
    ], size=12)
    s.notes_slide.notes_text_frame.text = (
        "Faces are classified by the sign of the normal's Y component after drafting (A releases +Y, B releases -Y) "
        "and by a straight-pull ray test (undercut). The PL is the set of mesh edges between A and B faces. "
        "Section rasters: air reachable from +Y = A steel, from -Y = B steel, enclosed = undercut.")

    # ---------------------------------------------------------------- 3 draft
    s = frame(prs, 3, "Draft analysis", "As received: every wall at 0°. DFM: 1° core side, 3° cavity side, 1° on the "
              "lifter-formed hook gap.", "DRAFT")
    image(s, F["draft"], 0.35, 1.45, 7.9, 5.5)
    da, db = inp["draft_area_mm2"], dfm["draft_area_mm2"]
    rows = [["Draft band", "As received", "DFM"],
            ["< 0.25°  (zero draft)", f"{da['0-0.25']:,.0f} mm²", f"{db['0-0.25']:.1f} mm²"],
            ["0.25° – 1°", f"{da['0.25-1.0']:.1f}", f"{db['0.25-1.0']:,.0f}"],
            ["1° – 3°", f"{da['1.0-3.0']:.1f}", f"{db['1.0-3.0']:.0f}"],
            ["3° – 80° (sloped)", f"{da['3.0-80.0']:.0f}", f"{db['3.0-80.0']:.0f}"],
            ["> 80° (flat to pull)", f"{da['80.0-90.01']:,.0f}", f"{db['80.0-90.01']:,.0f}"]]
    table(s, 8.5, 1.6, 4.35, rows, col_w=[1.85, 1.25, 1.25], size=11,
          colors={(1, 1): RED, (1, 2): GREEN})
    bullets(s, 8.5, 3.6, 4.35, 3.3, [
        ("1° core side:", "rails, lip, ribs, hook, latch, coring pins. The rib tip stays ≥ 0.93 mm."),
        ("3° cavity side:", "outer bumps, ready for texture (add 1° per 0.025 mm of texture depth)."),
        ("PL edges stay sharp.", "The latch catch face tilts 1° for release and keeps its 0.7 mm catch."),
        (f"Remaining {db['0-0.25']:.1f} mm²", "are the two hook-leg end faces at the lifter head. The lifter pulls away "
         "normal to them, so they need no draft."),
    ], size=12)
    s.notes_slide.notes_text_frame.text = (
        "Draft = asin(|n . Y|) per mesh face; 0.25-1 deg band in the DFM model is the 1 deg drafted walls measured "
        "on a 0.02 mm tessellation (they read 0.99-1.0 deg). Colours: red <0.25, yellow <0.9, light green <2.9, "
        "green sloped, grey flat to the pull.")

    # ---------------------------------------------------------------- 4 undercut
    uc = S["undercut"]
    lf = uc["lifter"]
    s = frame(prs, 4, "Undercut: retention hook",
              "The 1.0 mm hook gap cannot be pulled straight. The cover must stay closed, so the tool uses a lifter.",
              "UNDERCUT")
    image(s, F["undercut_sec"], 0.4, 1.5, 7.6, 2.55)
    image(s, F["undercut_3d"], 0.4, 4.15, 7.6, 2.8)
    rows = [["Option", "Tooling", "Part impact", "Verdict"],
            ["Angled lifter", f"{lf['angle_deg']:.0f}°, {lf['travel_mm']:.1f} mm travel in −Z → "
                              f"{lf['stroke_mm']:.0f} mm ejector stroke", "none: cover stays closed", "SELECTED"],
            ["Side slider", "gap opens toward the part interior (−Z)", "—", "not feasible"],
            ["Window through plate", "cavity steel through the plate", "hole in a battery cover: cells exposed, dust",
             "rejected"]]
    table(s, 8.2, 1.6, 4.65, rows, col_w=[1.1, 1.5, 1.0, 1.05], size=10, row_h=0.62,
          colors={(1, 3): GREEN, (2, 3): RED, (3, 3): RED})
    stat(s, 8.3, 4.5, 2.2, f"{uc['before_mm2']:.0f} mm²", "undercut, no release (as received)", color=RED, vsize=26)
    stat(s, 10.6, 4.5, 2.3, f"{uc['dfm_lifter_released_mm2']:.0f} mm²",
         f"released by the lifter  ·  {uc['dfm_unreleased_mm2']:.3f} mm² left", color=GREEN, vsize=26)
    text(s, 8.3, 5.75, 4.55, 1.2, [
        [("Lifter concept: ", {"bold": True}),
         ("L-shaped head fills the gap from the −Z side. The rod stands in front of the hook block, so it clears "
          "the block at 10°. The gap face on the block opens 1° toward −Z. Needs a mould base with ≥ 30 mm "
          "ejector stroke.", {})]], size=11, color=DARK)
    s.notes_slide.notes_text_frame.text = (
        "Undercut test: every face must see free space along its own release direction (+Y for A, -Y for B); "
        "vertical faces must escape either way. For the DFM part, each remaining undercut face is checked for free "
        "travel over the lifter stroke along -Z (relative motion of the lifter head). Lifter travel = gap depth "
        "4.5 mm + 0.5 mm clearance at 10 deg. A pass-through window would remove the lifter but puts a hole in a "
        "battery cover, so it was rejected.")

    # ---------------------------------------------------------------- 5 thickness
    ti, to = inp["thickness_mm"], dfm["thickness_mm"]
    s = frame(prs, 5, "Wall thickness",
              f"Nominal wall held at a uniform 2.0 mm; ribs at 50 % of wall. No thick nodes (max {to['max']:.2f} mm).",
              "THICKNESS")
    image(s, F["thickness"], 0.35, 1.45, 7.2, 3.3)
    cd = CategoryChartData()
    e = inp["thickness_hist"]["edges"]
    cd.categories = [f"{e[i]:.1f}–{e[i + 1]:.1f}" if e[i + 1] < 7 else f">{e[i]:.1f}" for i in range(len(e) - 1)]
    cd.add_series("As received", inp["thickness_hist"]["pct"])
    cd.add_series("DFM", dfm["thickness_hist"]["pct"])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(4.8), Inches(7.0), Inches(2.15), cd)
    ch = gf.chart
    style_chart(ch, 9)
    ch.has_title = True
    ch.chart_title.text_frame.text = "Share of surface by wall thickness [mm] (%)"
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    for ser, col in zip(ch.series, (GREY, AMBER)):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = col
    ch.value_axis.has_major_gridlines = True
    ch.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xE2, 0xE8)
    ch.value_axis.tick_labels.font.size = Pt(8)
    ch.category_axis.tick_labels.font.size = Pt(8)
    rows = [["Wall [mm]", "As received", "DFM"],
            ["min (0.5 %)", f"{ti['min']:.2f}", f"{to['min']:.2f}"],
            ["median", f"{ti['median']:.2f}", f"{to['median']:.2f}"],
            ["max (99.5 %)", f"{ti['max']:.2f}", f"{to['max']:.2f}"],
            ["rib root / wall", "0.8 / 2.0 (40 %)", "1.0 / 2.0 (50 %)"]]
    table(s, 7.9, 1.6, 4.95, rows, col_w=[1.55, 1.7, 1.7], size=11)
    bullets(s, 7.9, 3.35, 4.95, 3.6, [
        ("Primary wall 2.0 mm", "everywhere (plate + recess shell): inside the 2–3 mm target and uniform."),
        ("Ribs, rails and hook stay at 50–65 % of the wall",
         "so the class-A face does not show sink. Making them 2–3 mm would create 4 mm nodes."),
        ("Thinnest section 0.8 → 1.0 mm", "at the rib root. Measured minimum is the drafted tip."),
        ("Max 2.3 mm", "at radiused rib roots: +15 % only, so no local cooling problem."),
    ], size=12)
    s.notes_slide.notes_text_frame.text = (
        "Thickness = diameter of the largest inscribed sphere tangent at each of 30,000 surface samples, then a "
        "0.7 mm local max-filter to remove the known under-read next to sharp convex edges.")

    # ---------------------------------------------------------------- 6 gate
    G = S["gates"]
    pick = S["gate_pick"]
    s = frame(prs, 6, "Gate location & type",
              "Tunnel gate at the latch end: gate mark hidden, the snap arm fills first, no weld line on the cosmetic face.",
              "GATE")
    image(s, F["gates"], 0.4, 1.45, 8.4, 3.55)
    verdict = {"G1": ("hidden (top end)", "snap latch fills last", "reject"),
               "G2": ("visible side wall", "flow splits at recess", "reject"),
               "G3": ("hidden (end face)", "snap latch fills first", "SELECTED"),
               "G4": ("B face, 3-plate", "3-plate tool cost", "reject")}
    rows = [["Gate", "Type", "Max flow", "L / t", "Gate mark", "Fill behaviour", "Verdict"]]
    for k in ("G1", "G2", "G3", "G4"):
        rows.append([k, G[k]["label"].split(" - ")[1], f"{G[k]['max_flow_mm']:.0f} mm", f"{G[k]['L_over_t']:.0f}",
                     *verdict[k]])
    table(s, 0.5, 5.05, 8.3, rows, col_w=[0.55, 1.45, 0.95, 0.6, 1.55, 1.95, 1.25], size=10, row_h=0.36,
          colors={(3, 6): GREEN, (1, 6): RED, (2, 6): RED, (4, 6): RED})
    card(s, 9.05, 1.55, 3.8, 5.35)
    text(s, 9.25, 1.7, 3.4, 0.4, f"Selected: {pick}", size=16, font=HEAD, bold=True, color=AMBER)
    bullets(s, 9.25, 2.2, 3.45, 4.6, [
        ("Type:", "tunnel (submarine) gate from B, auto-degated on ejection. No manual trim."),
        ("Size:", "Ø1.0 mm tip (0.5 × wall), 40° tunnel, 1.2 mm land."),
        ("Where:", "bottom end face beside the thumb recess (X 3, Z 11.5), below the PL."),
        ("Flow:", f"L/t {G[pick]['L_over_t']:.0f}, far below the ABS limit of about {mc_limit(mat)}:1 at 2 mm."),
        ("Why not the others:", "G1 fills the snap latch last (short-shot and air-trap risk on the critical "
                                "feature). G2 leaves a visible gate mark. G4 needs a 3-plate tool."),
    ], size=12)
    s.notes_slide.notes_text_frame.text = (
        "Flow length = shortest surface path from the gate on a ~1.4 mm remeshed surface (Dijkstra). This is a "
        "geometric fill proxy, not a Moldflow result; confirm pressure/weld-line temperature in a fill simulation.")

    # ---------------------------------------------------------------- 7 ejection
    E = S["ejection"]
    n3 = sum(1 for p in E["pins"] if p["d"] >= 3)
    s = frame(prs, 7, "Ejection",
              f"{E['n_pins']} round pins on flat B-side faces. With 1° draft, the part releases without scuffing "
              "the ribs.", "EJECTION")
    image(s, F["ejection"], 0.4, 1.45, 8.4, 3.7)
    rows = [["#", "Location", "Face", "Pin"]]
    for i, p in enumerate(E["pins"], 1):
        rows.append([str(i), f"X {p['x']:.1f}  Z {p['z']:.1f}", p["on"], f"Ø{p['d']:.0f}"])
    table(s, 9.1, 1.6, 3.75, rows, col_w=[0.35, 1.35, 1.45, 0.6], size=9, row_h=0.27)
    stat(s, 0.6, 5.35, 2.4, f"{E['n_pins']} pins", f"{n3} × Ø3 plate · {E['n_pins'] - n3} × Ø2 hook / latch")
    stat(s, 3.3, 5.35, 2.4, f"{E['stroke_mm']:.0f} mm", f"stroke (core depth {E['core_depth_mm']:.0f} mm + 8)")
    bullets(s, 5.9, 5.3, 3.0, 1.6, [
        "Pins are checked to land fully on flat faces.",
        "Ø2 pin on the latch arm stops the catch from dragging.",
        "The lifter pushes the hook end, so no pin sits in its zone.",
    ], size=11)
    s.notes_slide.notes_text_frame.text = (
        "Each Ø3 pin site is verified by casting 17 rays over a disc of radius 2.0 mm (pin + 0.5 mm) and requiring "
        "all to land on the Y 18 inner face. Ribs (1.0 mm) are too thin for pins; add blade ejectors at rib ends "
        "only if trials show sticking. Witness marks are on the non-cosmetic inner face.")

    # ---------------------------------------------------------------- 8 cavities / tonnage
    T = S["tonnage"]
    cav = S["cavity_pick"]
    tp = next(r for r in T if r["cavities"] == cav)
    s = frame(prs, 8, "Cavitation & machine tonnage",
              f"{cav} cavities: {tp['clamp_calc_t']:.0f} t calculated clamp → {tp['machine_t']} t machine, "
              f"{tp['shot_g']:.0f} g shot, {tp['parts_per_h']} parts/h.", "CAVITIES")
    image(s, F["layout"], 0.4, 1.45, 4.6, 5.5)
    rows = [["Cavities", "Proj. area", "Clamp (calc)", "Machine", "Shot", "Parts / h"]]
    for r in T:
        rows.append([str(r["cavities"]), f"{r['proj_area_cm2']:.1f} cm²", f"{r['clamp_calc_t']:.1f} t",
                     f"{r['machine_t']} t", f"{r['shot_g']:.1f} g", f"{r['parts_per_h']:,}"])
    ci = [i for i, r in enumerate(T, 1) if r["cavities"] == cav][0]
    table(s, 5.3, 1.6, 7.55, rows, col_w=[1.0, 1.35, 1.45, 1.1, 1.15, 1.5], size=12, row_h=0.4,
          colors={(ci, j): AMBER for j in range(6)})
    card(s, 5.3, 3.85, 7.55, 1.2)
    text(s, 5.5, 3.95, 7.2, 1.0, [
        [("F = n · (A", {}), ("part", {"size": 9}), (" + A", {}), ("runner", {"size": 9}),
         (f") · p", {}), ("cavity", {"size": 9}), (" · SF", {})],
        [(f"A_part = {S['dfm']['reuse_checks_dfm']['projected_area_Y_cm2']:.2f} cm² (measured)  ·  A_runner ≈ 1.0 cm²/cavity  ·  "
          f"p = {mat['cavity_pressure_MPa']:.0f} MPa (ABS, 2 mm wall)  ·  SF = 1.2", {"size": 11, "color": GREY, "font": BODY})]],
        size=16, font=HEAD, color=DARK)
    bullets(s, 5.3, 5.25, 7.55, 1.7, [
        ("4 cavities", "suit typical remote-control volumes (0.3–1 M/yr at 85 % OEE). Move to 8 cavities on 100 t above ~1.5 M/yr."),
        ("Shot of 22 g", "should sit at 20–80 % of the barrel. A 50 t press with a Ø22–25 mm screw fits."),
    ], size=12)
    s.notes_slide.notes_text_frame.text = (
        "Part projected area is measured on the DFM STEP (sum of +Y-facing area x n_y). Runner allowance assumes "
        "Ø4/Ø3 H-balanced cold runner. Cycle time used for output: " + str(S["cooling"]["cycle"]["total_s"]) + " s. "
        "Annual volume was not supplied; the cavitation choice is an assumption to confirm with planning.")

    # ---------------------------------------------------------------- 9 cooling
    C, K = S["cooling"], S["coolant"]
    s = frame(prs, 9, "Cooling",
              f"Cooling time {C['t_cool_s']:.1f} s at the 2.0 mm wall; cycle ≈ {C['cycle']['total_s']:.0f} s. "
              "Separate A and B circuits held within 5 K.", "COOLING")
    image(s, F["cooling"], 0.4, 1.45, 8.5, 3.6)
    for i, (v, lab) in enumerate([(f"{C['t_cool_s']:.1f} s", "cooling time (2.0 mm)"),
                                  (f"{C['cycle']['total_s']:.0f} s", "cycle time"),
                                  (f"{K['flow_per_circuit_l_min']:.1f} L/min", "per circuit (Ø6)"),
                                  (f"Re {K['reynolds']:,}", "turbulent (> 10,000)")]):
        stat(s, 0.6 + i * 2.15, 5.25, 2.1, v, lab, vsize=24)
    card(s, 9.15, 1.55, 3.7, 5.35)
    bullets(s, 9.3, 1.7, 3.45, 5.1, [
        ("t = s²/(π²α) · ln[(4/π)(Tm−Tw)/(Te−Tw)]", f"s 2.0 mm, α {mat['diffusivity_mm2_s']} mm²/s, "
         f"Tm {mat['T_melt_C']:.0f}, Tw {mat['T_mold_C']:.0f}, Te {mat['T_eject_C']:.0f} °C."),
        ("Layout:", "Ø6 channels, 12 mm below the surface, 18 mm pitch, 2 per cavity per half, series circuit."),
        ("Latch core:", "Ø4 bubbler, where the 7 mm B-side step would otherwise run hot."),
        ("Heat load:", f"{K['heat_rate_W']} W → {K['flow_total_l_min']:.1f} L/min total for a 2 K water rise."),
        ("Thickest spot 2.3 mm", f"needs {C['t_cool_thickest_s']:.0f} s, but it is local: size the cycle on the 2.0 mm wall."),
    ], size=11)
    s.notes_slide.notes_text_frame.text = (
        "1-D plate cooling (mid-plane to ejection temperature). Heat per shot = m*cp*(Tm-Te) for the 4-cavity shot; "
        "water at 60 C (nu 0.47e-6 m2/s). Values are hand calculations to size the layout; confirm with a cooling "
        "analysis.")

    # ---------------------------------------------------------------- 10 fill
    FI = S["fill"]
    s = frame(prs, 10, "Fill sequence",
              f"From {FI['gate']}: latch and recess first, plate next, top lip and bumps last. "
              f"Max flow {FI['max_flow_mm']:.0f} mm.", "FILL")
    image(s, F["fill"], 0.35, 1.45, 8.4, 5.5)
    card(s, 9.0, 1.55, 3.85, 5.35)
    lf3 = FI["last_fill_xyz"]
    bullets(s, 9.15, 1.7, 3.6, 5.1, [
        ("0–20 %:", "snap arm and recess fill first, so orientation runs along the arm (strongest direction)."),
        ("20–80 %:", "the plate fills as a straight front; ribs fill along their length, with no hesitation."),
        ("80–100 %:", f"top lip, hook and bumps. Last to fill: X {lf3[0]:.0f}, Z {lf3[2]:.0f} (top corner)."),
        ("Vents:", "0.02 mm deep at the top corners, bump tips and the hook gap (vent along the lifter)."),
        ("Weld line:", "none on the cosmetic face (the plate has no openings). Fronts meet only around the hook block, on the inside."),
        ("Fill time ≈ 0.6 s", "then pack at 50–70 % of fill pressure."),
    ], size=11)
    s.notes_slide.notes_text_frame.text = (
        "Bands = percent of the maximum shortest surface path from the gate (geometric flow-front proxy). A flow "
        "simulation should confirm weld-line position and temperature, pressure drop and venting.")

    # ---------------------------------------------------------------- 11 warpage / deflection
    Wp = S["warpage"]
    s = frame(prs, 11, "Warpage & deflection",
              f"Bow ≈ {Wp['bow_per_K_mm']:.3f} mm per K of A/B mould ΔT. The source model's "
              f"{Wp['measured_source_bow_mm']:.2f} mm bow matches ≈ {Wp['dT_equivalent_K']:.0f} K.",
              "WARPAGE")
    cd = XyChartData()
    s1 = cd.add_series("Thermal bow, 65.5 mm plate")
    for dT in range(0, 21):
        s1.add_data_point(dT, round(Wp["bow_per_K_mm"] * dT, 4))
    s2 = cd.add_series("Measured in source model")
    s2.add_data_point(Wp["dT_equivalent_K"], Wp["measured_source_bow_mm"])
    s3 = cd.add_series("Target  ΔT ≤ 5 K")
    s3.add_data_point(5, round(Wp["bow_per_K_mm"] * 5, 4))
    gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Inches(0.5), Inches(1.5), Inches(6.6),
                            Inches(4.1), cd)
    ch = gf.chart
    style_chart(ch, 10)
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.series[0].format.line.color.rgb = STEEL
    ch.series[0].format.line.width = Pt(2.5)
    for ser, col in ((ch.series[1], RED), (ch.series[2], GREEN)):
        ser.format.line.fill.background()
        ser.marker.style = XL_MARKER_STYLE.CIRCLE
        ser.marker.size = 11
        ser.marker.format.fill.solid()
        ser.marker.format.fill.fore_color.rgb = col
        ser.marker.format.line.color.rgb = col
    va, xa = ch.value_axis, ch.category_axis
    va.has_title = True
    va.axis_title.text_frame.text = "Bow [mm]"
    xa.has_title = True
    xa.axis_title.text_frame.text = "A/B mould-temperature difference ΔT [K]"
    va.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xE2, 0xE8)
    va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
    xa.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
    image(s, F["section"], 0.5, 5.65, 6.6, 1.35)
    tb, ta = Wp["thumb_10N_before"], Wp["thumb_10N_after"]
    rows = [["Check", "As received", "DFM"],
            ["Bow at ΔT = 5 K", f"{Wp['bow_at_dT']['5']:.2f} mm", f"{Wp['bow_at_dT']['5']:.2f} mm"],
            ["Shrink over 71.5 mm", f"{Wp['shrink_mm'][0]:.2f}–{Wp['shrink_mm'][1]:.2f} mm", "same (tool +0.5 %)"],
            ["I (Z = 40 section)", f"{Wp['section_before']['I_mm4']:.1f} mm⁴", f"{Wp['section_after']['I_mm4']:.1f} mm⁴"],
            ["10 N thumb, mid-span", f"{tb['deflection_mm']:.2f} mm", f"{ta['deflection_mm']:.2f} mm"],
            ["Bending stress", f"{tb['stress_MPa']:.1f} MPa", f"{ta['stress_MPa']:.1f} MPa"],
            ["Safety vs yield", f"{tb['safety_vs_yield']:.1f}", f"{ta['safety_vs_yield']:.1f}"]]
    table(s, 7.45, 1.6, 5.4, rows, col_w=[2.0, 1.7, 1.7], size=11, row_h=0.36)
    bullets(s, 7.45, 4.3, 5.4, 2.6, [
        ("Keep ΔT(A−B) ≤ 5 K", f"to hold the bow ≤ {Wp['bow_at_dT']['5']:.2f} mm. Use separate A/B circuits and "
         "log the temperatures."),
        ("ABS is amorphous:", "shrink is near-isotropic, so the risk of differential-shrink warp is low."),
        ("Stiffness:", f"span {Wp['span_mm']:.0f} mm (latch to hook). Deflection stays below 0.4 mm and "
                       "stress sits at 1/5 of yield."),
    ], size=11)
    s.notes_slide.notes_text_frame.text = (
        "Thermal bow of a plate with a through-thickness mould-temperature difference: kappa = CTE*dT/t, "
        "delta = kappa*L^2/8 (L 65.5 mm plate, t 2.0 mm, CTE 90e-6/K). The 0.34 mm bow found in the reverse-"
        "engineered source mesh is consistent with about 14 K; this is an interpretation, not a measurement. "
        "Deflection: simply supported beam, 10 N at mid-span, E 2300 MPa, I from the measured Z=40 section.")

    # ---------------------------------------------------------------- 12 changes
    ch_ = S["dfm_changes"]
    s = frame(prs, 12, "DFM changes: before / after",
              "2-plate, 4-cavity tool with one lifter per cavity: closed cover, every wall drafted, every internal corner radiused.",
              "CHANGES")
    for i, (k, lab) in enumerate((("r_input_outer", "before · outer"), ("r_input_inner", "before · inner"),
                                  ("r_dfm_outer", "after · outer"), ("r_dfm_inner", "after · inner"))):
        x = 0.45 + i * 1.55
        card(s, x, 1.55, 1.45, 3.45, fill=LIGHT)
        image(s, F[k], x + 0.03, 1.6, 1.39, 3.1)
        text(s, x, 4.72, 1.45, 0.3, lab, size=10, color=RED if i < 2 else GREEN, align=PP_ALIGN.CENTER, bold=True)
    rows = [["Item", "Before", "After (part_DFM.step)"],
            ["Draft", "0° on all walls", f"{ch_['draft_core_deg']:.0f}° core · {ch_['draft_cavity_deg']:.0f}° cavity · "
                                          f"{ch_['draft_lifter_deg']:.0f}° lifter gap"],
            ["Hook undercut", f"{inp['undercut_area_mm2']:.0f} mm², no release defined", "closed hook kept; 10° lifter, gap face 1° drafted"],
            ["Internal radii", "R0 (sharp)", f"R{ch_['r_min_mm']} rib / rail / lip / recess / bump roots"],
            ["Thinnest section", f"rib {ch_['rib_w_before']} mm", f"rib {ch_['rib_w_after']} mm root (50 % t)"],
            ["Nominal wall", f"{ch_['wall_mm']} mm", f"{ch_['wall_mm']} mm (uniform)"],
            ["Outer bumps", "R2 barrel ends, 0° sides", "3° drafted pad, R0.75 top, R0.5 root"],
            ["Volume · mass", f"{inp['volume_mm3'] / 1000:.2f} cm³ · {inp['mass_g']:.2f} g",
             f"{dfm['volume_mm3'] / 1000:.2f} cm³ · {dfm['mass_g']:.2f} g"]]
    table(s, 6.75, 1.55, 6.1, rows, col_w=[1.45, 1.9, 2.75], size=10, row_h=0.37, bold_first_col=True)
    card(s, 0.45, 5.2, 12.4, 1.7)
    text(s, 0.65, 5.3, 12.0, 0.35, "Open items", size=14, font=HEAD, bold=True, color=AMBER)
    bullets(s, 0.65, 5.68, 5.9, 1.2, [
        ("Lifter:", "confirm ≥ 30 mm ejector stroke and rod clearance in the 4-cavity layout."),
        ("Rear case:", "check the 1.0 mm hook gap (now opening 1° toward −Z) against the housing edge."),
    ], size=11)
    bullets(s, 6.75, 5.68, 6.0, 1.2, [
        ("Simulation:", "run fill, pack and warp on the chosen ABS grade before tool release."),
        ("Snap-fit:", "run FEA of latch release strain with the rear-case geometry."),
    ], size=11)
    s.notes_slide.notes_text_frame.text = (
        "part_DFM.step is generated by battery_cover_RE/src/battery_cover_dfm.py from lib/geometry_dfm.py, which "
        "reuses the reverse-engineered lib/geometry.py (same frame, same recess and latch profiles).")

    prs.save(out)


def mc_limit(mat):
    return int(mat["max_flow_ratio"])
