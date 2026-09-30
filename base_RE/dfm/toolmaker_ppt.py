"""Toolmaker-style DFM report (python-pptx), the house format in .claude/skills/dfm-review/SKILL.md §4-5.

4:3 slides, "N.Topic" titles, green renders, red call-outs, CAV/COR indicator and an empty
RESULT/APPD/DATE sign-off table on every slide. Built from output/dfm_summary.json (run_dfm.py)."""

from __future__ import annotations

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.patches import Circle, FancyBboxPatch, Rectangle  # noqa: E402
from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

RED = RGBColor(0xFF, 0x00, 0x00)
BLACK = RGBColor(0, 0, 0)
GREY = RGBColor(0x80, 0x80, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"
W, H = 10.0, 7.5


# ============================================================ primitives
def txt(s, x, y, w, h, text, size=12, color=BLACK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        underline=False, italic=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Inches(0.03))
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        segs = ln if isinstance(ln, list) else ([ln] if isinstance(ln, tuple) else [(ln, {})])
        for seg, o in segs:
            r = p.add_run()
            r.text = seg
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.underline = o.get("underline", underline)
            f.color.rgb = o.get("color", color)
    return tb


def box(s, x, y, w, h, line=RED, width=1.5, dashed=False, fill=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = Pt(width)
    if dashed:
        sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    sh.shadow.inherit = False
    return sh


def label(s, x, y, text, size=11, w=None, h=0.27, bold=False, color=BLACK):
    w = w or max(0.45, 0.085 * len(text) * size / 11 + 0.16)
    b = box(s, x, y, w, h, fill=WHITE, width=1.25)
    tf = b.text_frame
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Inches(0.02))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return b


def comment(s, x, y, w, h, lines, size=11):
    b = box(s, x, y, w, h, fill=WHITE, width=1.5)
    tf = b.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Inches(0.08))
    lines = lines if isinstance(lines, list) else [lines]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        segs = ln if isinstance(ln, list) else ([ln] if isinstance(ln, tuple) else [(ln, {})])
        for seg, o in segs:
            r = p.add_run()
            r.text = seg
            r.font.name = FONT
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", False)
            r.font.color.rgb = o.get("color", BLACK)
    return b


def line(s, x1, y1, x2, y2, color=RED, width=1.5, arrow_end=False, arrow_start=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    for tag, on in (("a:tailEnd", arrow_end), ("a:headEnd", arrow_start)):
        if on:
            el = ln.makeelement(qn(tag), {"type": "triangle", "w": "med", "len": "med"})
            ln.append(el)
    return c


def arrow_block(s, x, y, w, h, direction="right"):
    shape = {"right": MSO_SHAPE.RIGHT_ARROW, "left": MSO_SHAPE.LEFT_ARROW, "up": MSO_SHAPE.UP_ARROW,
             "down": MSO_SHAPE.DOWN_ARROW}[direction]
    a = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = RED
    a.line.fill.background()
    return a


def pic(s, path, x, y, w=None, h=None):
    """Place an image inside the (x, y, w, h) box keeping its aspect; returns px -> slide-inch mapper."""
    im = Image.open(path)
    iw, ih = im.size
    if w and h:
        sc = min(w / iw, h / ih)
    elif w:
        sc = w / iw
    else:
        sc = h / ih
    pw, ph = iw * sc, ih * sc
    ox = x + ((w - pw) / 2 if (w and h) else 0)
    oy = y + ((h - ph) / 2 if (w and h) else 0)
    s.shapes.add_picture(str(path), Inches(ox), Inches(oy), Inches(pw), Inches(ph))
    return lambda px, py: (ox + px * sc, oy + py * sc), (ox, oy, pw, ph)


def crop_to_content(path, pad=12, tol=10):
    from PIL import ImageChops

    im = Image.open(path).convert("RGB")
    bg = Image.new("RGB", im.size, (255, 255, 255))
    diff = ImageChops.difference(im, bg).convert("L").point(lambda v: 255 if v > tol else 0)
    b = diff.getbbox()
    if not b:
        return path, (0, 0)
    b = (max(b[0] - pad, 0), max(b[1] - pad, 0), min(b[2] + pad, im.width), min(b[3] + pad, im.height))
    out = Path(path).with_name(Path(path).stem + "_c.png")
    im.crop(b).save(out)
    return out, (b[0], b[1])


def signoff(s):
    rows, cols = 3, 3
    t = s.shapes.add_table(rows, cols, Inches(0.15), Inches(6.62), Inches(2.1), Inches(0.66)).table
    t.columns[0].width = Inches(0.8)
    t.columns[1].width = Inches(0.65)
    t.columns[2].width = Inches(0.65)
    data = [["RESULT", "OK", "NG"], ["APPD", "", ""], ["DATE", "", ""]]
    for r in range(rows):
        t.rows[r].height = Inches(0.22)
        for c in range(cols):
            cell = t.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE
            cell.margin_left = cell.margin_right = Inches(0.04)
            cell.margin_top = cell.margin_bottom = Inches(0.0)
            tf = cell.text_frame
            tf.paragraphs[0].alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            r_ = tf.paragraphs[0].add_run()
            r_.text = data[r][c]
            r_.font.name = FONT
            r_.font.size = Pt(9)
            r_.font.color.rgb = BLACK
            tcPr = cell._tc.get_or_add_tcPr()
            for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = tcPr.makeelement(qn(tag), {"w": "6350"})
                sf = ln.makeelement(qn("a:solidFill"), {})
                clr = sf.makeelement(qn("a:srgbClr"), {"val": "A6A6A6"})
                sf.append(clr)
                ln.append(sf)
                tcPr.append(ln)
    tbl = t._tbl
    tblPr = tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        tblPr.remove(style)


def cavcor(s, x=9.05, y=0.1, pl=False):
    label(s, x, y, "CAV", w=0.55)
    label(s, x, y + 0.72, "COR", w=0.55)
    cx = x + 0.275
    line(s, cx, y + 0.27, cx, y + 0.72, arrow_end=True, arrow_start=True, width=1.25)
    line(s, x + 0.02, y + 0.495, x + 0.53, y + 0.495, width=1.25)
    if pl:
        label(s, x - 0.62, y + 0.37, "PL", w=0.45)


def new(prs, title=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if title:
        txt(s, 0.12, 0.08, 7.5, 0.4, title, size=16)
    signoff(s)
    return s


def note(s, text, y=7.22):
    txt(s, 2.4, y, 7.5, 0.25, text, size=8, color=GREY, italic=True)


# ============================================================ matplotlib helper figures
def fig_gate_layout(S, out):
    fig, ax = plt.subplots(figsize=(5.2, 6.0))
    L, Wp = 153.0, 51.0
    gap = 40.0
    for sx in (-1, 1):
        x0 = sx * gap / 2 - (Wp if sx < 0 else 0)
        ax.add_patch(Rectangle((x0, -L / 2), Wp, L, fc="#3DBE3D", ec="#1a5c1a", lw=1.2))
        ax.add_patch(Rectangle((x0 + 2, -L / 2 + 2), Wp - 4, L - 4, fc="#32A032", ec="#1a5c1a", lw=0.6))
        gz = 92.0 - L / 2
        gx = sx * gap / 2
        ax.plot([0, gx - sx * 4], [gz, gz], color="#7CFC7C", lw=6, solid_capstyle="round")
        ax.plot([gx - sx * 4, gx + sx * 2.6], [gz, gz], color="#9BFF9B", lw=2.2)
    ax.add_patch(Circle((0, 92.0 - L / 2), 3.2, fc="#7CFC7C", ec="#1a5c1a"))
    ax.text(0, 92.0 - L / 2 + 7, "sprue", ha="center", fontsize=9)
    ax.set_xlim(-100, 100)
    ax.set_ylim(-85, 85)
    ax.set_aspect("equal")
    ax.axis("off")
    fig.savefig(out, dpi=170, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def fig_gate_zoom(out):
    fig, ax = plt.subplots(figsize=(4.2, 3.6))
    ax.add_patch(Rectangle((0, 0), 2.0, 10, fc="#3DBE3D", ec="#1a5c1a"))          # side wall (section)
    ax.add_patch(Rectangle((-6, -1.2), 6, 1.2, fc="#C8C8C8", ec="k", lw=0.6))      # PL plate edge
    ax.plot([-9, -3.2], [-3.5, -3.5], color="#7CFC7C", lw=10, solid_capstyle="round")  # runner
    ax.plot([-3.2, 2.0], [-3.5, 1.6], color="#7CFC7C", lw=4.5)                     # tunnel
    ax.axhline(0, color="red", ls="--", lw=1)
    ax.text(-8.8, 0.25, "PL", color="red", fontsize=10)
    ax.text(2.4, 1.2, "Ø1.2 sub gate\n(inner wall face)", fontsize=9)
    ax.text(-8.8, -5.0, "Ø5.0 runner (core side)", fontsize=9)
    ax.set_xlim(-10, 8)
    ax.set_ylim(-6.5, 8)
    ax.set_aspect("equal")
    ax.axis("off")
    fig.savefig(out, dpi=170, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def fig_mold(S, out):
    M = S["mold"]
    fig, ax = plt.subplots(figsize=(5.4, 5.6))
    mw, ml, iw, il = M["mold_W_mm"], M["mold_L_mm"], M["insert_W_mm"], M["insert_L_mm"]
    ax.add_patch(Rectangle((-mw / 2 - 25, -ml / 2), mw + 50, ml, fc="#4F81BD", ec="k"))   # clamp plate
    ax.add_patch(Rectangle((-mw / 2, -ml / 2), mw, ml, fc="#FFE45C", ec="k"))              # A/B plates
    ax.add_patch(Rectangle((-iw / 2, -il / 2), iw, il, fc="#BFE3F5", ec="k"))              # insert
    for sx in (-1, 1):
        x0 = sx * 20 - (51 if sx < 0 else 0)
        ax.add_patch(Rectangle((x0, -76.5), 51, 153, fc="#3DBE3D", ec="#1a5c1a"))
    for cx in (-mw / 2 + 25, mw / 2 - 25):
        for cy in (-ml / 2 + 25, ml / 2 - 25):
            ax.add_patch(Circle((cx, cy), 10, fc="#DDDDDD", ec="k"))
    for yy in (-60, 0, 60):
        ax.plot([-mw / 2 - 25, mw / 2 + 25], [yy, yy], color="#1f77b4", lw=1.2, ls=":")
    ax.text(0, ml / 2 + 12, "Mold Top", ha="center", fontsize=10, color="red")
    ax.set_xlim(-mw / 2 - 45, mw / 2 + 45)
    ax.set_ylim(-ml / 2 - 25, ml / 2 + 30)
    ax.set_aspect("equal")
    ax.axis("off")
    fig.savefig(out, dpi=170, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def fig_legend(out, cmap, vmin, vmax, ticks, labelfmt="{:.2f}", title=None):
    fig, ax = plt.subplots(figsize=(0.9, 3.4))
    import numpy as np

    grad = np.linspace(vmax, vmin, 256)[:, None]
    ax.imshow(grad, aspect="auto", cmap=cmap, extent=[0, 1, vmin, vmax])
    ax.set_xticks([])
    ax.yaxis.tick_right()
    ax.set_yticks(ticks)
    ax.set_yticklabels([labelfmt.format(t) for t in ticks], fontsize=8)
    if title:
        ax.set_title(title, fontsize=8)
    fig.savefig(out, dpi=170, bbox_inches="tight", facecolor="white")
    plt.close(fig)


# ============================================================ deck
def build(S, repo: Path, out: Path):
    F = {k: repo / v for k, v in S["figures"].items()}
    FIG = out.parent / "figures"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    part, inp, dfm = S["part"], S["input"], S["dfm"]
    lift, ej, cool, tons = S["lifters"], S["ejection"], S["cooling"], S["tonnage"]
    pick = next(r for r in tons if r["cavities"] == S["cavity_pick"])
    n = [0]

    def title(t):
        n[0] += 1
        return f"{n[0]}.{t}"

    # ---------------------------------------------------------------- product / tooling information
    s = new(prs)
    txt(s, 0.3, 0.75, 3.8, 0.35, "Product information", size=13, color=RED, underline=True, align=PP_ALIGN.CENTER)
    info = [f"Part Name :- {part['name']}", f"Index No :- {part['code']}", f"Model :- {part['model']}",
            "Plastic material :- ABS  (assumed)", f"Shrinkage :- {part['shrinkage']:.3f}",
            f"Part Weight :- {dfm['mass_g']:.2f}g"]
    for i, t in enumerate(info):
        txt(s, 0.2, 1.3 + 0.48 * i, 4.2, 0.35, t, size=12)
    line(s, 4.55, 0.6, 4.55, 6.3, width=1.5)
    txt(s, 4.8, 0.75, 4.8, 0.35, "Tooling information", size=13, color=RED, underline=True, align=PP_ALIGN.CENTER)
    tool = ["Tooling type :-  2PLATE", "Hot Runner system :- NONE", "Gate type :- SUB GATE",
            f"No. of cavities :- 1 X {S['cavity_pick']}", "No. of slider :- NONE",
            f"No. of lifters :- {lift['count_per_cavity']}X / CAV", "Main insert material :- NAK-80",
            "Mold base material :- S50C", "Mold base Standard :- LKM",
            f"Injection machine capacity :-  {S['mold']['machine_t']}T"]
    for i, t in enumerate(tool):
        txt(s, 4.8, 1.3 + 0.48 * i, 5.0, 0.35, t, size=12)
    note(s, "Material not given on the TE-12XCME drawing: ABS assumed (same as the battery cover). Weight = DFM "
            "model volume x 1.05 g/cm3.", y=6.35)

    # ---------------------------------------------------------------- product size
    s = new(prs, "Product size")
    sz = S["size"]
    plan_c, off = crop_to_content(F["dfm_core_plan"])
    mp, (ox, oy, pw, ph) = pic(s, plan_c, 0.35, 0.8, 7.4, 3.4)
    side_c, off2 = crop_to_content(F["dfm_side"])
    mp2, (ox2, oy2, pw2, ph2) = pic(s, side_c, 0.35, 4.65, 7.4, 1.5)
    # length dimension under the plan, width at its right end, height at the side view
    line(s, ox, oy + ph + 0.12, ox + pw, oy + ph + 0.12, arrow_end=True, arrow_start=True)
    label(s, ox + pw / 2 - 0.4, oy + ph + 0.0, f"{sz['L']:.2f}", w=0.8)
    line(s, ox + pw + 0.15, oy, ox + pw + 0.15, oy + ph, arrow_end=True, arrow_start=True)
    label(s, ox + pw + 0.25, oy + ph / 2 - 0.13, f"{sz['W']:.2f}", w=0.75)
    line(s, ox2 + pw2 + 0.15, oy2, ox2 + pw2 + 0.15, oy2 + ph2, arrow_end=True, arrow_start=True)
    label(s, ox2 + pw2 + 0.25, oy2 + ph2 / 2 - 0.13, f"{sz['H']:.2f}", w=0.75)
    txt(s, 0.35, 4.3, 3, 0.3, "Core side (plan)", size=9, color=GREY)
    txt(s, 0.35, 6.2, 3, 0.3, "Side view (pull direction up)", size=9, color=GREY)

    # ---------------------------------------------------------------- parting line slides
    s = new(prs, title("Parting Line"))
    cavcor(s, 0.72, 0.5, pl=True)
    pic(s, crop_to_content(F["pl_side"])[0], 1.45, 0.62, 8.3, 1.2)
    pic(s, crop_to_content(F["pl_cav_iso"])[0], 0.3, 2.05, 4.6, 4.45)
    pic(s, crop_to_content(F["pl_core_iso"])[0], 5.0, 2.05, 4.8, 4.45)
    comment(s, 2.45, 6.62, 7.35, 0.62, ["Main PL flat at Y 4.34 (bottom of the side walls). PL steps down around the "
                                        "tongue at the top end and runs up the slanted open bottom end."], size=10)

    for key, ttl, zoom_key, txt_ in (
            ("tongue", "Parting Line", "pl_top_end", "Tongue below PL (Y 0.41-4.34) at the top end: PL steps down; "
                                                     "tongue formed in the core insert."),
            ("bottom", "Parting Line", "pl_bottom_end", "Open slanted bottom end (14.8 deg): end face on core side, "
                                                        "PL follows the outer edge of the end face."),
            ("opening", "Parting Line", "pl_opening", "Battery opening + cover slot: cavity / core shut-off along "
                                                      "the inner skin edge; rails and ledges on core side.")):
        s = new(prs, title(ttl))
        cavcor(s)
        base_img = F["pl_cav_iso"] if key == "opening" else F["pl_core_iso"]
        c, (cx0, cy0) = crop_to_content(base_img)
        mp, _ = pic(s, c, 0.25, 0.55, 4.6, 5.9)
        px = S["px"]["pl_cav_iso" if key == "opening" else "pl_core_iso"][key][0]
        bx, by = mp(px[0] - cx0, px[1] - cy0)
        box(s, bx - 0.35, by - 0.3, 0.7, 0.6)
        zc = crop_to_content(F[zoom_key])[0]
        _, (zx, zy, zw, zh) = pic(s, zc, 5.05, 1.2, 4.75, 4.6)
        box(s, zx, zy, zw, zh, dashed=True, width=1.0)
        line(s, bx + 0.35, by, zx, zy + zh / 2, width=1.0)
        comment(s, 2.45, 6.62, 7.35, 0.62, txt_, size=10)

    # ---------------------------------------------------------------- lifters
    s = new(prs, title("Lifter"))
    c, (cx0, cy0) = crop_to_content(F["lifter_plan"])
    mp, (ox, oy, pw, ph) = pic(s, c, 0.3, 0.9, 9.4, 4.3)
    teeth = S["px"]["lifter_plan"]["teeth"]
    k = len(teeth) // 2
    for i, (px_, py_) in enumerate(teeth):
        x, y = mp(px_ - cx0, py_ - cy0)
        left = i < k  # left-hand teeth (X 2.62) sit on one long edge of the plan
        lab = f"Lifter {i % k + 1:02d}" + ("L" if left else "R")
        ly = y - 0.62 if y < oy + ph / 2 else y + 0.35
        label(s, x - 0.42, ly, lab, size=9, w=0.84, h=0.24)
        d = "down" if y < oy + ph / 2 else "up"
        arrow_block(s, x - 0.08, y + (0.08 if d == "down" else -0.36), 0.16, 0.28, d)
    comment(s, 2.45, 5.55, 7.35, 0.95,
            [f"{lift['count_per_cavity']} lifters / cavity: one per PL catch tooth "
             f"({TOOTH_TXT(lift)}).",
             f"Lifter travel {lift['travel_mm']:.1f} mm inward, angle {lift['calc']['angle_deg']:.0f} deg -> "
             f"ejector stroke {lift['calc']['ejector_stroke_mm']:.1f} mm min."], size=10)

    s = new(prs, title("Lifter"))
    c, (cx0, cy0) = crop_to_content(F["lifter_zoom"])
    pic(s, c, 0.25, 0.6, 4.4, 4.3)
    label(s, 1.6, 0.62, "Lifter 04L", size=10, w=1.0)
    pic(s, F["sec_tooth"], 4.8, 0.7, 5.0, 3.0)
    line(s, 4.75, 0.6, 4.75, 5.0, width=1.25)
    rel = lift["released_within_travel_mm2"]
    comment(s, 4.9, 3.85, 4.9, 1.45,
            [f"Undercut faces in lifter zones: {lift['undercut_in_lifter_zones_mm2']:.1f} mm2",
             f"Released within {lift['travel_mm']:.1f} mm travel: {rel:.1f} mm2 "
             f"(not released: {lift['not_released_mm2']:.2f} mm2)",
             "Window through side wall rejected: cosmetic side wall + weak snap"], size=10)
    note(s, "Ray check on the DFM model: every lifter-zone undercut face is tested along the lifter travel "
            "direction (+X left side, -X right side).", y=5.4)

    # ---------------------------------------------------------------- ejector
    s = new(prs, title("Ejector"))
    c, (cx0, cy0) = crop_to_content(F["eject_plan"])
    pic(s, c, 0.25, 0.9, 7.2, 4.6)
    from collections import Counter

    cnt = Counter(p["d"] for p in ej["pins"])
    y0 = 1.0
    colors = {1.5: RGBColor(0x33, 0x33, 0xFF), 2.0: RGBColor(0xFF, 0x00, 0xFF), 2.5: RGBColor(0xFF, 0x99, 0x00),
              3.0: RGBColor(0xFF, 0x99, 0x00), 3.5: RGBColor(0xFF, 0x66, 0x00), 4.0: RGBColor(0xCC, 0x33, 0x00)}
    for d, k_ in sorted(cnt.items()):
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.6), Inches(y0), Inches(0.1 + 0.05 * d),
                                 Inches(0.1 + 0.05 * d))
        dot.fill.solid()
        dot.fill.fore_color.rgb = RGBColor(0xDD, 0x22, 0x22)
        dot.line.fill.background()
        label(s, 8.1, y0 - 0.02, f"EP {d:.1f} ( {k_}X )", w=1.6)
        y0 += 0.5
    bl = Counter((b["w"], b["l"]) for b in ej["blades"])
    for (w_, l_), k_ in sorted(bl.items()):
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.55), Inches(y0 + 0.05), Inches(0.42), Inches(0.14))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor(0x1A, 0x40, 0xE6)
        bar.line.fill.background()
        label(s, 8.1, y0 - 0.02, f"EB {l_:.1f} X {w_:.1f} ( {k_}X )", w=1.8, size=10)
        y0 += 0.45
        if y0 > 5.6:
            break
    comment(s, 2.45, 5.75, 7.35, 0.75,
            [f"Pins on flat core-side faces (disc check), none in lifter zones. Core depth {ej['core_depth_mm']:.1f} mm, "
             f"ejector stroke {ej['stroke_mm']:.1f} mm."], size=10)

    # ---------------------------------------------------------------- gate
    s = new(prs, title("Gate"))
    lay = FIG / "gate_layout.png"
    fig_gate_layout(S, lay)
    pic(s, lay, 0.3, 0.7, 4.3, 5.7)
    zoom = FIG / "gate_zoom.png"
    fig_gate_zoom(zoom)
    _, (zx, zy, zw, zh) = pic(s, zoom, 4.9, 1.45, 4.85, 3.9)
    box(s, zx, zy, zw, zh, dashed=True, width=1.0)
    label(s, 6.9, 0.55, "Cold Runner + Sub Gate", w=2.6, size=12)
    label(s, 5.9, 1.0, "Ø1.2", w=0.7)
    G = S["gates"]
    rows = [f"{k}: {v['label']} - max flow {v['max_flow_mm']:.0f} mm (L/t {v['L_over_t']:.0f})" for k, v in G.items()]
    txt(s, 4.9, 5.4, 4.9, 1.0, rows, size=8)
    comment(s, 2.45, 6.62, 7.35, 0.62, f"Gate {S['gate_pick']} chosen: no mark on the cosmetic skin, shortest "
                                       f"practical flow {G[S['gate_pick']]['max_flow_mm']:.0f} mm, away from "
                                       f"lifters.", size=10)

    # ---------------------------------------------------------------- draft
    leg = FIG / "legend_draft.png"
    fig_legend(leg, matplotlib.colors.LinearSegmentedColormap.from_list("d", ["#2649F2", "#3DBE3D", "#D91AD9"]),
               -0.5, 0.5, [0.5, 0.3333, 0.1667, 0, -0.1667, -0.3333, -0.5], "{:+.4f}")
    for key, head, msg in (
            ("input", "Draft angle analysis", [("All Green Surface Are 0º Propose To Add 1.0º For Cavity Side & "
                                                "Draft 0.5º For Core Side", {})]),
            ("dfm", "Draft angle analysis (DFM)", [("After DFM: side walls 1.0º cav / 0.5º core, thin ribs 0.5º. "
                                                    "Remaining green = block/detail walls: add in tool design", {})])):
        s = new(prs, title(head))
        pic(s, crop_to_content(F[f"draft_{key}_cav"])[0], 0.25, 0.6, 4.3, 5.2)
        pic(s, crop_to_content(F[f"draft_{key}_core"])[0], 4.6, 0.6, 4.4, 5.2)
        pic(s, leg, 9.1, 0.8, 0.8, 3.0)
        zd = S[key]["draft_area_mm2"]["0-0.25"]
        comment(s, 2.45, 5.95, 7.35, 0.6, msg, size=10)
        txt(s, 2.45, 6.58, 7.3, 0.3, f"0º area {zd:,.0f} mm2 (of {S[key]['area_mm2']:,.0f} mm2)", size=9,
            color=RED, bold=True)

    # ---------------------------------------------------------------- thickness
    s = new(prs, title("Thickness mark"))
    legt = FIG / "legend_thick.png"
    fig_legend(legt, "jet", 0, 3.5, [0, 0.5, 1, 1.5, 2, 2.5, 3, 3.5], "{:.1f}", "mm")
    pic(s, crop_to_content(F["thick_cav"])[0], 0.25, 0.6, 4.3, 5.3)
    pic(s, crop_to_content(F["thick_core"])[0], 4.6, 0.6, 4.4, 5.3)
    pic(s, legt, 9.1, 0.8, 0.8, 3.0)
    t = dfm["thickness_mm"]
    comment(s, 2.45, 5.95, 7.35, 0.62, [f"Nominal wall {t['median']:.2f} mm, ribs 0.8 mm. Red surface have thickness "
                                        f"mark: contact walls Z 141-146 up to 3.0 mm."], size=10)

    # ---------------------------------------------------------------- suggest slides
    s = new(prs, title("Suggest"))
    cavcor(s)
    c, (cx0, cy0) = crop_to_content(F["undercut_core_iso"])
    mp, _ = pic(s, c, 0.25, 0.55, 4.3, 5.4)
    label(s, 1.5, 0.58, "Undercut", w=1.0)
    pic(s, F["sec_tooth"], 4.7, 1.05, 5.1, 2.8)
    comment(s, 4.75, 3.95, 5.05, 1.9,
            ["Have undercut: 10 PL catch teeth (1.0 x 1.5 x 4.0) for the front-case snap.",
             "Function part, keep shape. Suggest to make undercut with lifter (10X / CAV).",
             "Window in side wall rejected (cosmetic + weak snap)."], size=10)

    s = new(prs, title("Suggest"))
    c, (cx0, cy0) = crop_to_content(F["steel_core"])
    pic(s, c, 0.25, 0.55, 3.6, 5.3)
    label(s, 1.2, 0.58, "Weak steel", w=1.1)
    pic(s, F["sec_steel"], 3.95, 0.55, 5.85, 2.75)
    pic(s, F["sec_steel2"], 3.95, 3.3, 5.85, 2.55)
    label(s, 3.95, 3.0, "Before", w=0.8)
    label(s, 9.0, 3.0, "After", w=0.7)
    comment(s, 2.45, 5.95, 7.35, 0.62,
            [f"Weak steel 0-0.75 mm between side plates/ribs and R8 fillet, easy broken. Suggest to add plastic "
             f"(fill gap), {S['wedges']['count']} places, slot below kept."], size=10)

    s = new(prs, title("Suggest"))
    pic(s, F["sec_rib"], 0.3, 0.7, 9.4, 4.3)
    label(s, 0.5, 1.25, "Before", w=0.8)
    label(s, 5.2, 1.25, "After", w=0.7)
    rb = S["ribs"]
    we = rb["wall_example"]
    comment(s, 2.45, 5.2, 7.35, 1.2,
            ["Ribs 0.8-1.0 mm (0.4-0.5 x wall): no thickness mark, but 0º draft -> sticking on core, ejection marks.",
             f"Suggest 0.5º draft each side, tip keep (0.8 / 1.0 mm), root grows: this wall {we['tip_mm']:.1f} -> "
             f"{we['root_mm']:.2f} mm over {we['height_mm']:.1f} mm. Plates at side-wall slots and rims at cover "
             "rails keep 0º (function)."], size=10)

    s = new(prs, title("Suggest"))
    pic(s, F["sec_thick"], 0.3, 0.6, 4.6, 4.8)
    c, (cx0, cy0) = crop_to_content(F["thick_core"])
    pic(s, c, 5.0, 0.6, 4.8, 4.8)
    comment(s, 2.45, 5.55, 7.35, 0.95,
            [f"Contact walls Z 141-146: up to 3.0 mm (1.5 x wall) -> thickness mark risk and "
             f"{cool['t_cool_3mm_spots_s']:.0f} s cooling vs {cool['t_cool_nominal_2mm_s']:.0f} s.",
             "Suggest to core out to 2.0 mm if the battery contact allows (customer to confirm; not changed in STEP)."],
            size=10)

    s = new(prs, title("Suggest"))
    pic(s, crop_to_content(F["draft_dfm_core"])[0], 0.25, 0.6, 4.4, 5.0)
    zd = dfm["zero_draft_by_region_mm2"]
    rows = [f"{k}: {v:,.0f} mm2" for k, v in list(zd.items())[:6]]
    txt(s, 4.9, 0.8, 4.9, 2.5, ["Remaining 0º faces after DFM (by area):"] + rows, size=10)
    comment(s, 4.8, 3.3, 5.0, 2.3,
            ["Add draft in tool design: blocks, battery bay walls, end walls, detail walls: 0.5º core / 1.0º cavity.",
             "Keep 0º on: cover rails & rims (cover fit), catch faces (snap).",
             "Add R0.3-R0.5 at rib roots (electrode radius)."], size=10)
    note(s, "The DFM STEP adds the draft on the parametric walls and the thin ribs; the section-driven detail "
            "prisms keep their measured 0º faces.", y=5.75)

    # ---------------------------------------------------------------- mold layout
    s = new(prs, title("Mold Layout"))
    ml = FIG / "mold_layout.png"
    fig_mold(S, ml)
    pic(s, ml, 0.4, 0.6, 5.2, 5.8)
    M = S["mold"]
    label(s, 5.9, 0.9, f"Mold {M['mold_W_mm']:.0f} X {M['mold_L_mm']:.0f}", w=2.2)
    label(s, 5.9, 1.4, f"Insert {M['insert_W_mm']} X {M['insert_L_mm']}", w=2.2)
    label(s, 5.9, 1.9, f"{M['machine_t']} TON", w=1.2, bold=True)
    if M["machine_t"] != M["machine_t_clamp"]:
        comment(s, 5.9, 5.0, 3.9, 0.9, [[(f"Machine size need to change from {M['machine_t_clamp']}T to "
                                          f"{M['machine_t']}T", {"bold": True, "color": RED})],
                                        f"(mold {max(M['mold_W_mm'], M['mold_L_mm']):.0f} mm > "
                                        f"{M['tie_bar_clamp_machine_mm']} mm tie bars)"], size=10)
    rows = [f"{r['cavities']} cav: clamp {r['clamp_calc_t']:.0f} t -> {r['machine_t']} T, shot {r['shot_g']:.0f} g, "
            f"{r['parts_per_h']} pcs/h" for r in tons]
    txt(s, 5.9, 2.5, 3.9, 1.6, rows, size=9)
    txt(s, 5.9, 3.7, 3.9, 1.3, [f"Projected area {dfm['projected_area_cm2']:.1f} cm2/cav, cavity pressure "
                                 f"{S['material']['cavity_pressure_MPa']:.0f} MPa, SF 1.2 (hand calc).",
                                 M["tie_bar_note"]], size=9)

    # ---------------------------------------------------------------- simulation proxies
    legr = FIG / "legend_fill.png"
    fig_legend(legr, "jet", 0, S["fill"]["fill_time_s_est"], [0, S["fill"]["fill_time_s_est"] / 2,
                                                              S["fill"]["fill_time_s_est"]], "{:.2f}", "s")
    s = new(prs, title("Fill Time"))
    pic(s, crop_to_content(F["fill_core"])[0], 0.25, 0.6, 4.3, 5.2)
    pic(s, crop_to_content(F["fill_cav"])[0], 4.6, 0.6, 4.4, 5.2)
    pic(s, legr, 9.1, 0.8, 0.8, 3.0)
    comment(s, 2.45, 5.95, 7.35, 0.55, [(f"Filling time = {S['fill']['fill_time_s_est']:.2f}sec",
                                          {"bold": True, "color": RED, "size": 14})])
    note(s, "estimate, geometric proxy, not Moldflow: flow-length map from the gate; time = shot volume / 40 cm3/s.",
         y=6.55)

    s = new(prs, title("Ejection Temperature"))
    lege = FIG / "legend_ej.png"
    fig_legend(lege, "jet", 0, 22, [0, 5, 10, 15, 20], "{:.0f}", "s")
    pic(s, crop_to_content(F["ejtime_core"])[0], 0.25, 0.6, 5.2, 5.2)
    pic(s, lege, 5.6, 0.8, 0.8, 3.0)
    comment(s, 6.5, 1.0, 3.3, 2.2, [f"Time to reach ejection temperature (85 °C): median "
                                    f"{S['ejection_time_s']['p50']:.1f} s, p99 {S['ejection_time_s']['p99']:.1f} s.",
                                    "Slowest: 3.0 mm contact walls."], size=10)
    note(s, "estimate, 1-D plate cooling per local thickness, not Moldflow.", y=6.55)

    s = new(prs, title("Air Traps"))
    pic(s, crop_to_content(F["air_core"])[0], 0.25, 0.6, 4.3, 5.2)
    pic(s, crop_to_content(F["air_cav"])[0], 4.6, 0.6, 4.4, 5.2)
    comment(s, 2.45, 5.95, 7.35, 0.55, [("Make air vent around parts (last-fill points: red)",
                                          {"bold": True, "color": RED, "size": 12})])
    note(s, "estimate, geometric proxy, not Moldflow: local maxima of flow length from the gate.", y=6.55)

    s = new(prs, title("Weld Lines"))
    pic(s, crop_to_content(F["weld_cav"])[0], 0.25, 0.6, 4.3, 5.2)
    pic(s, crop_to_content(F["weld_core"])[0], 4.6, 0.6, 4.4, 5.2)
    comment(s, 2.45, 5.95, 7.35, 0.55, [("Weld lines where flow fronts meet (red); check cosmetic side at the "
                                          "battery opening ends", {"bold": True, "color": RED, "size": 11})])
    note(s, "estimate, geometric proxy, not Moldflow: edges where the flow-front directions oppose.", y=6.55)

    # ---------------------------------------------------------------- cycle time
    s = new(prs, title("Cycle Time"))
    cy = cool["cycle"]
    cd = CategoryChartData()
    cd.categories = ["Fill time", "Pack time", "Cooling time", "Mold open time"]
    cd.add_series("s", (cy["fill_s"], cy["pack_s"], cy["cool_s"], cy["mold_open_s"]))
    gf = s.shapes.add_chart(XL_CHART_TYPE.PIE, Inches(0.3), Inches(0.7), Inches(4.6), Inches(3.8), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(9)
    pl_ = ch.plots[0]
    pl_.has_data_labels = True
    pl_.data_labels.number_format = '0.0"s"'
    pl_.data_labels.number_format_is_linked = False
    pl_.data_labels.font.size = Pt(9)
    pl_.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    to_3d_pie(ch)  # last: python-pptx cannot read the chart back once it is 3-D
    pic(s, crop_to_content(F["dfm_cav_iso"])[0], 5.2, 0.6, 4.6, 2.6)
    lines_ = [f"1.Fill time = {cy['fill_s']:.3f}sec", f"2.Pack time = {cy['pack_s']:.1f} sec",
              f"3.Cooling time = {cy['cool_s']:.1f} sec", f"4.Mold open time = {cy['mold_open_s']:.1f} sec"]
    txt(s, 5.2, 3.35, 4.6, 1.3, lines_, size=12)
    txt(s, 0.4, 4.75, 9.4, 0.4, [[("Cycle time = Fill + Pack + Cooling + Mold open = ", {}),
                                  (f"{cy['total_s']:.3f} sec", {"color": RED, "underline": True, "bold": True})]],
        size=14)
    note(s, f"Cooling calculated (1-D plate, ABS, {cool['governing_wall_mm']:.2f} mm = 99.5 % of the part); "
            "fill = estimate; pack 8.0 s and mold open 5.0 s = supplier defaults from the reference reports.", y=5.3)

    prs.save(out)
    return out


def to_3d_pie(chart):
    """python-pptx has no 3-D pie writer: convert the pie chart XML to c:pie3DChart with a tilted view."""
    cs = chart._chartSpace
    c = cs.find(qn("c:chart"))
    pie = c.find(qn("c:plotArea")).find(qn("c:pieChart"))
    pie.tag = qn("c:pie3DChart")
    fsa = pie.find(qn("c:firstSliceAng"))
    if fsa is not None:
        pie.remove(fsa)
    v3 = c.makeelement(qn("c:view3D"), {})
    for tag, val in (("c:rotX", "30"), ("c:rotY", "20"), ("c:rAngAx", "0")):
        v3.append(v3.makeelement(qn(tag), {"val": val}))
    c.insert(list(c).index(c.find(qn("c:plotArea"))), v3)


def TOOTH_TXT(lift):
    t = lift["tooth"]
    return f"{t['protrusion']:.1f} x {t['height']:.1f} x {t['length']:.1f} mm, catch face Y {t['catch_face_Y']}"
